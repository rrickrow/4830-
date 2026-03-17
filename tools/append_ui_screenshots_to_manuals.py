from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "软著申请材料_拆分版"
SHOT_DIR = DOC_DIR / "运行截图"


MANUALS = [
    DOC_DIR / "方案一_松辽流域河道植被遥感分析软件V1.0_操作手册.docx",
    DOC_DIR / "方案二_流域生态监测成果展示与自动报告平台V1.0_操作手册.docx",
]


def set_default_font(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(11)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    set_default_font(p)


def append_section(doc_path: Path) -> None:
    doc = Document(doc_path)
    doc.add_page_break()

    doc.add_heading("八、真实运行界面截图说明", level=1)
    add_body(
        doc,
        "以下截图均通过实际运行当前软件后，以本地无头浏览器访问系统页面并自动截取而成，不属于示意性占位图片。与前文中插入的分析结果图相比，本节所展示的内容更强调软件在真实运行状态下的界面结构、交互区域和成果展示方式，可用于补充说明软件已经形成完整的可视化交互界面。",
    )
    add_body(
        doc,
        "从软著材料编写角度看，这类界面截图能够更直接地证明软件具备可操作的界面入口、参数配置能力和运行结果展示能力。对于审查者而言，界面截图与结果图配合出现，能够同时体现“可以运行”和“有输出结果”两个层面的软件成型度。",
    )

    screenshots = [
        (
            "01_首页参数区.png",
            "图7 真实运行截图：系统首页与参数配置区域",
            "该截图展示了软件启动后的首页界面。页面顶部集中提供研究区域、年份范围、年份步长、数据源和报告选项等参数输入控件，用户可在该区域完成一次任务的核心配置。此类界面体现了软件具备明确的交互入口和参数化运行能力。",
        ),
        (
            "02_河道变化页.png",
            "图8 真实运行截图：河道变化分析页面",
            "该截图展示了软件完成计算后默认进入的河道变化页面。界面中包含年度水体掩膜图、年度指标趋势图、中心线叠加图、交互式地图及河道指标表，能够从图像、地图和表格三个层面同步展示河道变化结果，说明软件已形成较完整的结果展示页面。",
        ),
        (
            "03_植被响应页.png",
            "图9 真实运行截图：植被响应分析页面",
            "该截图展示了植被响应页面的真实运行状态。页面中集中呈现缓冲带FVC热力图、河道变化与植被响应耦合散点图以及缓冲带统计表，能够较好反映软件对植被响应结果的组织能力和可视化表达能力。",
        ),
        (
            "04_驱动评估页.png",
            "图10 真实运行截图：驱动评估页面",
            "该截图展示了驱动评估页面的实际显示效果。页面中包含标准化回归系数图、系数表和模型指标信息，便于用户从界面层面直接读取统计分析结果，也表明软件具备将分析结果转化为可解释可展示内容的能力。",
        ),
        (
            "05_报告导出页.png",
            "图11 真实运行截图：报告导出页面",
            "该截图展示了报告导出页面。该页面集中列出本次运行生成的各类文件，并为用户提供下载入口。此类页面能够清晰体现软件的成果归档与文件分发能力，是软件从计算走向交付的重要组成部分。",
        ),
    ]

    for filename, caption, desc in screenshots:
        image_path = SHOT_DIR / filename
        if not image_path.exists():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(image_path), width=Cm(15.6))
        add_caption(doc, caption)
        add_body(doc, desc)

    doc.save(doc_path)


def main() -> None:
    for path in MANUALS:
        append_section(path)


if __name__ == "__main__":
    main()
