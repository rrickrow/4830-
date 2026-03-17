from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from river_insight.domain.models import AnalysisRequest
from river_insight.services.analysis_service import AnalysisService


DOC_DIR = ROOT / "软著申请材料_拆分版"


@dataclass(frozen=True)
class ManualSpec:
    title: str
    short_name: str
    filename: str
    study_area: str
    positioning: str
    target_user: str
    goal_text: str
    emphasis_text: str


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if style_name == "Title":
            style.font.size = Pt(20)
        elif style_name == "Heading 1":
            style.font.size = Pt(15)
        else:
            style.font.size = Pt(12)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(subtitle)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(12)
    doc.add_paragraph()


def add_paragraph(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_image(doc: Document, image_path: Path, caption: str, width_cm: float = 15.5) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp_run = cp.add_run(caption)
    cp_run.font.name = "宋体"
    cp_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    cp_run.font.size = Pt(10.5)


def run_analysis(study_area: str):
    service = AnalysisService()
    request = AnalysisRequest(
        study_area=study_area,
        start_year=1995,
        end_year=2025,
        year_step=10,
        provider="demo",
        include_report=True,
    )
    return service.run(request)


def build_manual(spec: ManualSpec, result) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc, f"{spec.title}V1.0 操作手册", f"简称：{spec.short_name}")

    doc.add_heading("一、软件概述", level=1)
    add_paragraph(
        doc,
        f"{spec.title}是依托“松辽流域河道变迁与植被响应原型系统”拆分形成的独立软件成果。该软件围绕{spec.positioning}展开，以本地部署、本地分析和本地导出的方式提供完整的软件使用链路。用户在普通Windows环境下即可完成参数配置、分析执行、结果查看和成果导出，无需额外搭建前后端分离环境，也不依赖Node生态工具链。",
    )
    add_paragraph(
        doc,
        "从软件著作权申报角度看，本软件已经具备清晰的软件名称、明确的功能边界、稳定的程序结构和可复现的运行结果。相较于原始项目中的整体原型，本手册所对应的软件版本强调独立的软件能力表达，便于在申请材料中突出其作为独立应用软件的业务价值与技术完整性。",
    )

    doc.add_heading("二、建设目标与适用对象", level=1)
    add_paragraph(doc, spec.goal_text)
    add_paragraph(doc, spec.target_user)
    add_paragraph(
        doc,
        spec.emphasis_text,
    )

    doc.add_heading("三、运行环境与启动方式", level=1)
    add_paragraph(
        doc,
        "推荐在Windows 10或Windows 11 64位系统中运行本软件。硬件方面建议使用双核及以上处理器、8GB及以上内存和足够的本地磁盘空间，以确保图表渲染、空间结果导出和文档生成过程保持良好响应。开发与调试环境采用Python 3.13.5，正式运行环境建议保持Python 3.11及以上版本。",
    )
    add_paragraph(
        doc,
        "运行前需要安装项目依赖，包括Streamlit、Pandas、GeoPandas、Shapely、NumPy、SciPy、scikit-image、Statsmodels、Folium、Jinja2和ReportLab等组件。完成依赖安装后，用户通过命令“streamlit run streamlit_app.py”即可启动本地服务，并在浏览器中访问系统界面。首次启动后，系统会根据默认参数直接支持演示数据路径的完整运行。",
    )
    add_paragraph(
        doc,
        "本软件当前版本默认使用内置演示数据，能够在未配置真实GEE凭据的情况下完成全部核心流程。若用户后续需要接入真实遥感数据，可在保留当前界面和流程结构的前提下，进一步扩展GEE接口实现，从而实现更贴近生产场景的数据分析能力。",
    )

    doc.add_heading("四、软件整体使用流程", level=1)
    add_paragraph(
        doc,
        "软件的标准使用流程可以概括为“参数配置—自动分析—结果浏览—成果导出”四个阶段。用户首先在首页填写研究区域、分析年份范围、年份步长、数据源和报告选项；系统随后自动完成数据解析、遥感指标计算、河道提取、植被统计和驱动评估；分析完成后，用户分别在不同标签页查看图表、表格和空间结果；最终可在导出页面下载本次任务生成的全部成果文件。",
    )
    add_paragraph(
        doc,
        "这种流程将传统上分散在脚本、GIS软件、图表工具和文档工具中的步骤收敛到同一套软件交互中，显著降低了重复操作成本。对于课程演示、答辩展示或项目阶段性汇报，这种统一化的使用方式能够有效提升结果展示的规范性与效率。",
    )

    doc.add_heading("五、主要功能介绍与详细操作说明", level=1)

    doc.add_heading("5.1 首页参数配置与任务创建", level=2)
    add_paragraph(
        doc,
        "进入软件后，用户首先看到的是参数配置区域。该区域集中提供研究区域、开始年份、结束年份、年份步长、数据源和报告开关等核心参数。所有参数均采用显式配置方式，能够让使用者在任务开始前清楚掌握当前分析的对象、时间范围和输出要求。",
    )
    add_paragraph(
        doc,
        "在填写参数时，软件会进行基本校验。例如，若开始年份大于结束年份，界面将直接阻止提交；若年份范围超出演示数据支持区间，软件同样会给出明确提示。这种前置校验机制可以减少无效计算，提高任务执行的稳定性，并降低因参数错误导致的结果偏差风险。",
    )
    add_paragraph(
        doc,
        "当用户点击“开始分析”后，系统会自动调用内部分析服务，按预设顺序执行数据准备、指标计算、空间分析、图表生成、结果导出和报告生成等步骤。对于普通使用者而言，无需逐个调用底层脚本，仅通过统一按钮即可完成任务调度。",
    )

    doc.add_heading("5.2 河道变化分析", level=2)
    add_paragraph(
        doc,
        "河道变化分析页面用于展示不同年份下的水体范围、河道结构和关键指标结果。系统会基于NDWI指数执行水体识别，并在形态学清理与连通域过滤后得到较稳定的河道掩膜，再进一步通过骨架化方法提取中心线和关键节点，从而形成可量化的河道变化结果。",
    )
    add_paragraph(
        doc,
        "该页面通常包含年度水体掩膜图、年度指标趋势图、中心线叠加图和河道指标表。通过这些内容，用户既可以从图像角度理解河道范围变化，也可以从指标角度读取河道面积、中心线长度、平均河宽、迁移量和摆动幅度等关键数值，为后续分析提供直观依据。",
    )
    add_image(doc, Path(result.artifacts["water_masks_panel_png"]), "图1 软件运行结果示意图：年度水体掩膜图")
    add_paragraph(
        doc,
        "上图为软件在实际运行中生成的年度水体掩膜图。该图能够直观展示不同年份的河道水体空间形态变化，是识别河道边界变化与水域收缩、扩张趋势的重要可视化成果。对于论文图件、阶段汇报和教学演示而言，该结果具有较高的展示价值。",
    )
    add_image(doc, Path(result.artifacts["centerline_overlay_png"]), "图2 软件运行结果示意图：河道中心线与节点叠加图")
    add_paragraph(
        doc,
        "上图展示了河道中心线及节点叠加结果。中心线能够反映河槽主干路径，节点则反映交汇或端点位置。通过不同年份的中心线比较，用户可更清晰地观察河道迁移方向及空间摆动特征，这类结果尤其适合用作河道形态演化分析的辅助依据。",
    )
    add_image(doc, Path(result.artifacts["annual_metrics_png"]), "图3 软件运行结果示意图：年度河道指标趋势图")
    add_paragraph(
        doc,
        "年度河道指标趋势图将水体面积和平均河宽等结果以时间序列方式表达出来。与单幅图像相比，趋势图更适合用于快速把握长期变化规律，并帮助使用者在短时间内识别异常年份、波动区间和总体演化方向。",
    )

    doc.add_heading("5.3 植被响应分析", level=2)
    add_paragraph(
        doc,
        "植被响应分析页面围绕NDVI和FVC两个核心指标展开。系统首先计算像元层面的植被指数，再结合河道水体结果计算河岸距离，并按预设缓冲区分带统计不同距离区间内的植被覆盖情况。这一过程能够将河道变化与植被变化建立起较直接的空间联系。",
    )
    add_paragraph(
        doc,
        "该页面的重点在于回答“河道变化是否在不同距离上对植被产生不同影响”这一问题。用户可通过热力图和统计表快速判断近河带、中间带和远河带在FVC均值及变化率上的差异，为生态响应分析提供更细粒度的判断依据。",
    )
    add_image(doc, Path(result.artifacts["fvc_heatmap_png"]), "图4 软件运行结果示意图：缓冲带FVC热力图")
    add_paragraph(
        doc,
        "上图为缓冲带FVC热力图。不同颜色深浅反映不同缓冲带在各年份下的植被覆盖水平，能够帮助使用者快速比较各空间距离层中的植被变化差异。该结果尤其适合用于说明河道变化对周边植被的空间层次影响。",
    )
    add_image(doc, Path(result.artifacts["coupling_png"]), "图5 软件运行结果示意图：河道变化与植被响应耦合散点图")
    add_paragraph(
        doc,
        "耦合散点图用于展示河宽变化或河道迁移与FVC变化之间的关系。通过散点分布，使用者可在正式回归分析之前先观察变量之间是否存在趋势性关系，这对解释后续统计结果和撰写研究结论都具有重要帮助。",
    )

    doc.add_heading("5.4 驱动评估与统计分析", level=2)
    add_paragraph(
        doc,
        "驱动评估页面面向更偏分析解释层的使用需求。系统会根据河道变化指标、植被变化指标及辅助驱动变量构建回归样本，并对变量进行标准化处理后执行OLS回归，从而得到各驱动因素对植被变化的影响系数、显著性水平和模型拟合度指标。",
    )
    add_paragraph(
        doc,
        "该页面不仅适合技术人员进行结果解读，也适合在汇报场景中向非技术受众说明不同因素的影响方向。例如，当某一因子的标准化回归系数明显为正或为负时，使用者可以较直观地说明该因子对植被变化的促进或抑制作用。",
    )
    add_image(doc, Path(result.artifacts["regression_png"]), "图6 软件运行结果示意图：标准化回归系数图")
    add_paragraph(
        doc,
        "上图为标准化回归系数图。该图以图形方式表达各变量的相对影响强度，便于在汇报材料中快速展示回归结果核心结论。相比直接阅读系数表，图形化表达更适合对外展示和答辩场景。",
    )

    doc.add_heading("5.5 成果导出与自动报告", level=2)
    add_paragraph(
        doc,
        "在完成分析后，软件会自动生成带时间戳的独立输出目录，并将本次任务产生的CSV、GeoJSON、PNG和PDF等文件统一保存。用户可在导出页面中查看每个产物的名称，并根据需要逐项下载。这种输出机制有助于保持任务批次清晰，便于日后回溯与归档。",
    )
    add_paragraph(
        doc,
        "若用户勾选了生成报告选项，系统会在分析结束后自动将参数摘要、关键图表、年度指标和回归结果整理为PDF报告。该文档可直接用于课程汇报、项目答辩或内部资料留存，减少人工整理图表和重新排版的工作量。",
    )
    add_paragraph(
        doc,
        "对于软件著作权申报材料而言，自动报告和统一导出能力能够体现软件已经形成较完整的业务闭环，即不仅能够完成内部计算，也能够稳定输出可直接使用的成果文档，这有助于增强软件的应用属性和成品属性表达。",
    )

    doc.add_heading("六、异常处理与使用建议", level=1)
    add_paragraph(
        doc,
        "当前版本已经包含基础异常处理机制。若参数超出范围、样本数量不足或数据源不可用，系统会给出可读提示并尽量避免程序直接中断。对于首次使用者，建议优先使用默认演示数据完成一次完整运行，以熟悉软件流程和结果组织方式；在确认流程稳定后，再根据后续需求扩展到更复杂的数据接入场景。",
    )
    add_paragraph(
        doc,
        "在实际撰写申报和答辩材料时，建议优先选取本手册中插入的这些真实运行结果图片作为软件运行截图。这些图像均由当前软件在实际执行过程中自动生成，能够较好地证明软件具备实际运行能力和稳定输出能力，也更符合软著操作手册对“运行示意”的常见表达方式。",
    )

    doc.add_heading("七、版本说明", level=1)
    add_paragraph(
        doc,
        "本手册对应的软件版本为V1.0。本版本已具备相对完整的软件结构、稳定的运行流程和清晰的结果输出能力，能够作为独立软件成果进行说明和申报。后续若需要进一步提升两套软著材料之间的区分度，可继续在现有基础上扩展界面流程、调整模块范围或增加更具差异化的业务功能。",
    )

    output_path = DOC_DIR / spec.filename
    doc.save(output_path)
    return output_path


def main() -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)

    suite_a = ManualSpec(
        title="松辽流域河道植被遥感分析软件",
        short_name="河道植被分析软件",
        filename="方案一_松辽流域河道植被遥感分析软件V1.0_操作手册.docx",
        study_area="songhua",
        positioning="河道与植被变化分析引擎",
        target_user="本软件主要面向遥感测绘、水文地理、生态环境监测和高校科研教学等场景，适用于需要开展河道演化识别、植被变化分析和驱动诊断的教师、学生、科研人员及项目研发成员。",
        goal_text="本软件的建设目标在于将河道变化识别、植被响应统计和多因子评估等关键能力整合成一套可独立运行的分析型软件，使研究人员能够用更低的使用门槛完成从计算到解释的核心分析流程。软件强调分析深度与量化能力，适合作为偏算法与偏研究表达的软件成果进行申报。",
        emphasis_text="在拆分申报口径上，本软件突出的是“分析引擎”属性。也就是说，软件更强调对河道与植被变化进行计算、识别和量化评估的能力，而不是以平台化展示为主，这一点与另一套侧重展示与自动报告的平台型软件形成清晰区分。",
    )

    suite_b = ManualSpec(
        title="流域生态监测成果展示与自动报告平台",
        short_name="生态监测报告平台",
        filename="方案二_流域生态监测成果展示与自动报告平台V1.0_操作手册.docx",
        study_area="combined",
        positioning="成果展示、任务编排与自动报告平台",
        target_user="本软件主要面向生态监测展示、科研汇报、教学演示和项目交付等场景，适用于需要快速查看分析结果、统一下载成果文件并自动生成正式报告的教师、学生、团队负责人和项目管理者。",
        goal_text="本软件的建设目标在于将分析任务配置、结果展示、图形表达、空间文件下载和报告自动生成整合到同一平台界面中，使使用者能够在最短路径内完成从任务创建到成果归档的全过程。软件更强调交互组织、成果展示和交付效率，适合作为平台型应用软件进行申报。",
        emphasis_text="在拆分申报口径上，本软件突出的是“监测展示与自动报告平台”属性。与偏分析引擎的软件相比，它更强调统一界面、任务编排、图表呈现、成果下载和文档生成等能力，从而在功能定位与使用场景上形成明确差异。",
    )

    build_manual(suite_a, run_analysis(suite_a.study_area))
    build_manual(suite_b, run_analysis(suite_b.study_area))


if __name__ == "__main__":
    main()
