from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from river_insight.domain.models import AnalysisRequest
from river_insight.services.analysis_service import AnalysisService


DOC_DIR = ROOT / "软著申请材料_拆分版"
SHOT_DIR = DOC_DIR / "运行截图"


@dataclass(frozen=True)
class ManualSpec:
    title: str
    short_name: str
    output_name: str
    study_area: str
    positioning: str
    vision: str
    users: str
    strengths: str
    process_text: str
    tech_text: str
    env_text: str
    interface_notes: list[tuple[str, str, str, str]]


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if style_name == "Title":
            style.font.size = Pt(20)
        elif style_name == "Heading 1":
            style.font.size = Pt(15)
        else:
            style.font.size = Pt(12)


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sp.add_run(subtitle)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(12)
    doc.add_paragraph()


def add_body(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_center_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)


def add_image(doc: Document, image_path: Path, caption: str, width_cm: float = 15.4) -> None:
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    add_center_caption(doc, caption)


def run_analysis(study_area: str):
    service = AnalysisService()
    request = AnalysisRequest(
        study_area=study_area,
        start_year=1995,
        end_year=2025,
        year_step=10,
        provider="demo",
        include_report=True,
    )
    return service.run(request)


def build_manual(spec: ManualSpec, result) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc, f"{spec.title} V1.0 操作手册", f"简称：{spec.short_name}")

    doc.add_heading("一、软件概述", level=1)
    add_body(
        doc,
        f"{spec.title}是依托“松辽流域河道变迁与植被响应原型系统”拆分形成的独立软件成果，当前版本重点面向{spec.positioning}。软件采用本地部署方式运行，通过统一的图形界面组织任务配置、分析执行、结果浏览和成果导出，使用户能够在单一软件中完成完整操作流程，而不需要借助多个分散工具协同处理。",
    )
    add_body(
        doc,
        "与面向演示的原型系统相比，本手册对应的软件版本更加突出软件边界的独立性和操作路径的完整性。其核心价值不仅体现在具备可执行的业务逻辑，也体现在具备明确的界面入口、可重复的操作过程和稳定的成果输出能力，这些内容均可直接用于软件著作权申报材料中的软件说明部分。",
    )

    doc.add_heading("二、产品愿景与目标", level=1)
    doc.add_heading("2.1. 产品愿景", level=2)
    add_body(doc, spec.vision)
    doc.add_heading("2.2. 目标行业与用户群体", level=2)
    add_body(doc, spec.users)

    doc.add_heading("三、软件特色与优势", level=1)
    add_body(doc, spec.strengths)
    add_image(
        doc,
        Path(result.artifacts["regression_png"]),
        "软件特色示意图（以系统生成的统计分析图为例）",
        width_cm=13.8,
    )
    add_body(
        doc,
        "上图为软件在真实运行中自动生成的统计分析图形。该类图形并非静态示意，而是系统根据当前任务数据自动输出的结果文件。其能够直观体现软件在可视化表达、图表组织和分析结果呈现方面的成型度，也是本软件区别于单纯脚本程序的重要特征之一。",
    )

    doc.add_heading("四、使用流程简述", level=1)
    add_body(doc, spec.process_text)
    add_image(doc, SHOT_DIR / "01_首页参数区.png", "使用流程示意图（首页参数配置界面）", width_cm=11.5)
    add_body(
        doc,
        "从实际操作顺序看，用户通常先在首页完成参数配置并启动任务，随后进入河道变化、植被响应和驱动评估等页面查看结果，最后在报告导出页面下载文件并完成归档。该流程在界面层面清晰可见，符合模板中强调的“操作路径简洁、步骤清楚、便于上手”的写法要求。",
    )

    doc.add_heading("五、技术特点", level=1)
    add_body(doc, spec.tech_text)
    add_image(
        doc,
        Path(result.artifacts["centerline_overlay_png"]),
        "技术架构示意图（以系统生成的空间分析结果为例）",
        width_cm=12.8,
    )
    add_body(
        doc,
        "该图反映了软件在空间分析过程中的实际输出能力。虽然本软件不以复杂架构图作为主要展示形式，但通过真实的空间结果图，可以直接说明系统已经具备数据处理、空间计算和成果渲染的完整技术链条。",
    )

    doc.add_heading("六、系统运行环境", level=1)
    add_body(doc, spec.env_text)
    add_body(
        doc,
        "当前版本可在Windows 10或Windows 11 64位系统中稳定运行，建议使用8GB及以上内存的个人电脑执行分析任务。系统依赖Python 3.11及以上版本以及Streamlit、Pandas、GeoPandas、NumPy、SciPy、scikit-image、Statsmodels、Folium、Jinja2和ReportLab等组件，依赖安装完成后即可直接启动本地服务。",
    )

    doc.add_heading("七、主要功能介绍与操作说明", level=1)
    for index, (name, shot_name, intro, operation) in enumerate(spec.interface_notes, start=1):
        doc.add_heading(f"{index}. {name}", level=2)
        add_image(doc, SHOT_DIR / shot_name, f"{name}", width_cm=11.8)
        doc.add_heading(f"{index}.1. 功能简介", level=3)
        add_body(doc, intro)
        doc.add_heading(f"{index}.2. 操作说明", level=3)
        add_body(doc, operation)
        add_body(
            doc,
            "在实际使用时，建议用户先观察当前页面顶部或主体区域的主要控件，再按照“先配置、后执行、再查看、最后导出”的顺序进行操作。若界面中出现提示信息或等待状态，应优先等待系统完成当前步骤后再切换页面，以避免在分析尚未结束时中断结果加载。",
        )

    path = DOC_DIR / spec.output_name
    doc.save(path)
    return path


def main() -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)

    scheme_a = ManualSpec(
        title="松辽流域河道植被遥感分析软件",
        short_name="河道植被分析软件",
        output_name="方案一_松辽流域河道植被遥感分析软件V1.0_操作手册.docx",
        study_area="songhua",
        positioning="河道与植被变化分析",
        vision="本软件的产品愿景是围绕流域遥感分析场景，构建一套可直接运行、可重复操作、可稳定输出结果的分析型应用软件。软件希望在传统遥感研究依赖多工具切换和人工整理的基础上，进一步降低操作门槛，使使用者能够通过更清晰的界面流程完成河道变化识别、植被响应统计和驱动因素评估等核心任务，从而提升科研分析与成果整理效率。",
        users="本软件主要面向遥感测绘、水文地理、生态环境监测、高校科研训练和课程教学等场景，适用于需要开展河道结构识别、植被覆盖分析、生态响应判断及驱动因素量化的教师、学生、科研人员和项目研发成员。对于需要进行阶段汇报或快速形成分析结果的团队而言，本软件也具备较好的应用价值。",
        strengths="本软件的优势在于其分析链路完整、操作方式统一且结果表达直观。软件既能执行多时相数据准备、NDWI/NDVI/FVC计算、河道中心线识别和多因子回归分析等核心逻辑，又能将结果整理为图表、表格和文档进行展示与导出。与纯脚本工具相比，软件在界面组织和操作流程上更便于非程序背景用户使用；与传统分散式处理方式相比，软件在效率和结果一致性方面具有更明显优势。",
        process_text="用户在启动软件后，首先需要在首页完成研究区域、分析年份、年份步长、数据源和报告选项等参数设置。参数确认无误后，点击“开始分析”按钮，系统将自动进入计算过程，并依次完成数据准备、指标计算、河道提取、植被统计和驱动分析。分析完成后，用户可通过不同页面查看河道变化结果、植被响应结果和统计评估结果，最后在报告导出页面下载本次任务的全部成果文件。",
        tech_text="本软件在技术上采用模块化Python架构，将数据模型、数据提供、预处理、指标计算、空间分析、统计分析和结果导出划分为相互协作的独立模块。界面层使用Streamlit构建本地Web交互页面，空间分析部分使用GeoPandas、Shapely、scikit-image和SciPy实现河道识别、中心线提取和距离统计，统计部分使用Statsmodels完成OLS回归。技术路线兼顾可维护性、可扩展性和实际可运行性，适合作为分析型软件成果进行独立申报。",
        env_text="本软件建议运行在Windows 10或Windows 11 64位系统环境中。硬件建议为双核及以上处理器、8GB及以上内存和足够的本地磁盘空间。运行前需安装Python 3.11及以上版本，并按照依赖清单安装Streamlit、Pandas、GeoPandas、NumPy、SciPy、scikit-image、Statsmodels、Folium、Jinja2和ReportLab等支持组件。软件当前版本支持在未配置GEE凭据的情况下通过内置演示数据直接完成全流程运行。",
        interface_notes=[
            (
                "参数配置与启动界面",
                "01_首页参数区.png",
                "参数配置与启动界面是用户进入系统后的第一操作界面，承担分析任务创建和输入约束校验的作用。通过这一界面，用户可以明确设置研究对象、分析时间范围、年份步长、数据源类型及是否生成报告，从而确定本次分析任务的执行边界。",
                "打开系统后，首先在页面顶部依次选择研究区域、开始年份、结束年份、年份步长和数据源。若需要在分析结束后同步生成报告，可勾选“生成报告”选项。确认参数无误后，点击“开始分析”按钮，系统即开始执行任务。若输入参数不符合要求，界面会直接给出提示，用户应先修改参数后再重新发起任务。",
            ),
            (
                "河道变化界面",
                "02_河道变化页.png",
                "河道变化界面用于集中展示河道水体范围、中心线结构和年度指标结果。该界面是本软件分析链路中的核心结果页之一，能够帮助用户直观查看不同年份下河道形态的变化过程，并结合表格数据读取关键指标值。",
                "当系统分析完成后，默认可先查看河道变化界面。用户可自上而下浏览年度水体掩膜图、中心线叠加图和年度趋势图，并在页面下方查看河道指标表。若需要进一步比较年份差异，可结合图形结果和表格中的水体面积、中心线长度、平均河宽等字段进行交叉判断。",
            ),
            (
                "植被响应界面",
                "03_植被响应页.png",
                "植被响应界面用于展示不同河岸缓冲带中的植被覆盖变化情况。该界面通过热力图、耦合散点图和统计表共同表达植被变化在空间距离和时间维度上的差异，是解释河道变化如何影响植被的重要界面。",
                "在界面顶部标签中点击“植被响应”即可进入该页面。进入后，用户可先查看左侧热力图，快速识别不同缓冲区在各年份下的FVC变化；随后查看右侧散点图，判断河道变化与植被变化之间是否存在趋势关系；最后结合下方统计表读取各缓冲带的详细数值，用于支撑分析结论。",
            ),
            (
                "驱动评估界面",
                "04_驱动评估页.png",
                "驱动评估界面用于展示多因子回归分析结果。该界面将回归系数图、系数明细表和模型指标集中呈现，使用户能够快速识别不同驱动因素对植被变化的影响方向和影响强度。",
                "点击顶部“驱动评估”标签后进入该界面。用户首先查看系数图，初步判断各变量的影响方向；随后阅读下方系数表，重点关注系数值、显著性和区间范围；最后结合页面中的样本数、R²和调整后R²等信息，对模型的解释能力进行判断。此页面适合用于形成研究结论或汇报摘要。",
            ),
            (
                "报告导出界面",
                "05_报告导出页.png",
                "报告导出界面是本软件的结果出口页面。该界面集中列出本次任务产生的全部成果文件，并提供下载入口，使用户能够快速完成结果归档和对外分发。",
                "点击顶部“报告导出”标签后，页面会显示本次运行对应的输出目录及全部产物列表。用户可逐项点击下载按钮，获取PDF报告、CSV指标文件、GeoJSON空间数据和图片文件。若用于课程汇报或项目申报，建议优先下载PDF报告与关键图表，再根据需要补充原始数据文件。",
            ),
        ],
    )

    scheme_b = ManualSpec(
        title="流域生态监测成果展示与自动报告平台",
        short_name="生态监测报告平台",
        output_name="方案二_流域生态监测成果展示与自动报告平台V1.0_操作手册.docx",
        study_area="combined",
        positioning="成果展示、交互浏览与自动报告生成",
        vision="本软件的产品愿景是将流域生态监测场景中的参数配置、结果展示、图表浏览、成果下载和文档归档整合到一个统一平台中，使用户能够通过更短的操作路径完成从任务创建到成果输出的完整过程。软件强调界面化表达和交付便利性，希望在教学演示、项目汇报和成果归档等场景中提供更标准化的软件支持。",
        users="本软件主要面向生态监测展示、科研汇报、教学答辩和项目成果交付场景，适用于需要快速浏览分析结果、统一下载文件并自动生成汇报文档的教师、学生、项目管理者和研究团队成员。对于更重视“看结果、讲结果、交结果”的使用者而言，本软件具有更高的适用性。",
        strengths="本软件最大的优势在于平台化组织能力强、界面操作路径清晰、成果导出过程完整。软件将多类分析结果分配到不同页面，并通过统一的标签页进行组织，使用户可以在同一界面中快速切换不同结果视图。与此同时，系统还能自动整理输出目录并生成PDF报告，使其非常适合用于展示、汇报和归档等面向交付的场景。",
        process_text="用户在首页完成参数设置并点击“开始分析”后，系统会自动执行后台分析流程。分析完成后，用户可依次进入河道变化、植被响应和驱动评估等页面浏览结果，并在最后的报告导出页面统一下载成果文件。整个过程遵循“先配置、后生成、再查看、最后导出”的平台化操作逻辑，能够显著简化结果整理工作。",
        tech_text="本软件采用“界面层—分析服务层—导出服务层—报告服务层”的分层结构进行构建。界面层负责参数输入和页面展示，分析服务层负责任务编排与数据处理，导出服务层负责生成结构化文件，报告服务层负责将结果整理为PDF文档。技术上使用Streamlit、Folium和Matplotlib构建结果展示能力，并结合Python数据处理生态完成后端逻辑，是典型的面向展示和交付的软件平台实现方式。",
        env_text="本软件建议部署于Windows 10或Windows 11 64位环境。为保证界面渲染和文档导出流畅，建议使用8GB及以上内存的个人电脑。运行环境依赖Python 3.11及以上版本以及Streamlit、Pandas、GeoPandas、NumPy、SciPy、scikit-image、Statsmodels、Folium、Jinja2和ReportLab等组件。完成依赖安装后，可直接通过命令启动本地服务，无需单独配置复杂的前后端工程环境。",
        interface_notes=[
            (
                "首页参数配置界面",
                "01_首页参数区.png",
                "首页参数配置界面是整个平台的任务入口。平台在该界面中集中放置所有核心配置项，使用户在任务启动前即可清楚看到本次分析涉及的区域、时间和输出方式。该设计符合平台型软件强调统一入口和统一操作逻辑的需求。",
                "进入系统后，用户首先在首页填写研究区域、开始年份、结束年份和年份步长，并选择数据源类型。若需要输出报告，则保持“生成报告”选项勾选状态。参数确认后，点击“开始分析”按钮即可启动任务。平台会在后台自动处理数据，并在分析完成后刷新各结果页面内容。",
            ),
            (
                "河道变化展示界面",
                "02_河道变化页.png",
                "河道变化展示界面是平台默认结果页，用于承载最直观的时空变化信息。通过图形、地图和指标表的组合，平台可以将复杂的分析结果转化为更适合展示和汇报的页面内容，便于用户快速获取整体认知。",
                "任务完成后，用户可直接在该页面查看水体掩膜、河道趋势图、中心线叠加图和河道指标表。建议先浏览上方图形结果，再阅读下方指标表，以形成“先看趋势、再看数值”的阅读顺序。若用于答辩展示，该页面通常可作为第一个重点说明页面。",
            ),
            (
                "植被响应展示界面",
                "03_植被响应页.png",
                "植被响应展示界面承担的是结果解释功能，重点展示不同缓冲区的植被变化情况以及河道变化与植被变化之间的关系。对于汇报场景而言，该界面比单纯的数据表更容易帮助听众理解生态响应结果。",
                "点击“植被响应”标签进入页面后，用户可依次查看热力图、耦合散点图和缓冲带统计表。建议先说明热力图中的整体变化趋势，再通过散点图解释变量关系，最后用统计表补充具体数值。这样的讲解顺序与模板中常见的“界面截图后接操作说明”写法较为一致。",
            ),
            (
                "驱动评估展示界面",
                "04_驱动评估页.png",
                "驱动评估展示界面用于承载统计分析结果，是平台中最适合用于形成汇报结论的页面之一。页面结构清晰，系数图、系数表和模型指标集中呈现，能够帮助用户快速组织表达重点。",
                "进入该页面后，用户可先观察标准化回归系数图，识别哪些变量的影响更明显；再查看下方表格中的具体系数和显著性；最后读取样本数与拟合度等信息，说明模型的可靠程度。对于写报告或准备答辩材料，这一页面的结果通常可以直接转化为正文内容。",
            ),
            (
                "报告导出界面",
                "05_报告导出页.png",
                "报告导出界面是平台中最能体现“自动交付”价值的功能页面。平台在该页面中统一列出本次分析的所有文件，使用户无需离开系统即可完成成果下载和归档。",
                "用户点击“报告导出”标签后，可在页面中查看输出目录及各类产物列表。根据实际需求，可优先下载PDF报告用于汇报展示，也可下载CSV、GeoJSON和图片文件用于二次加工。该页面清楚体现了平台从分析到交付的闭环能力，是本软件区别于纯分析型软件的重要特点。",
            ),
        ],
    )

    build_manual(scheme_a, run_analysis(scheme_a.study_area))
    build_manual(scheme_b, run_analysis(scheme_b.study_area))


if __name__ == "__main__":
    main()
