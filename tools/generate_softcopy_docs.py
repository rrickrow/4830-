from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "软著申请材料_拆分版"


@dataclass(frozen=True)
class SuiteSpec:
    key: str
    full_name: str
    short_name: str
    version: str
    positioning: str
    purpose: str
    audience: str
    main_functions: list[str]
    technical_features: list[str]
    operation_sections: list[tuple[str, str]]
    code_files: list[Path]


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if style_name == "Title":
            style.font.size = Pt(20)
        elif style_name == "Heading 1":
            style.font.size = Pt(15)
        else:
            style.font.size = Pt(12)

    if "CodeBlock" not in doc.styles:
        style = doc.styles.add_style("CodeBlock", 1)
        style.font.name = "Consolas"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(8)


def add_title(doc: Document, title: str, subtitle: str | None = None) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle)
        sub_run.font.name = "微软雅黑"
        sub_run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        sub_run.font.size = Pt(12)
    doc.add_paragraph()


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"{label}：")
    run.bold = True
    p.add_run(text)


def add_body_paragraph(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def build_application_doc(spec: SuiteSpec) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc, f"{spec.full_name}{spec.version} 申请材料（申请表）", "软件著作权申报草稿")

    doc.add_heading("一、软件著作权信息", level=1)
    add_labeled_paragraph(doc, "软件全称", spec.full_name)
    add_labeled_paragraph(doc, "软件简称", spec.short_name)
    add_labeled_paragraph(doc, "版本号", spec.version)
    add_labeled_paragraph(doc, "权利取得方式", "原始取得")
    add_labeled_paragraph(doc, "权利范围", "全部权利")
    add_labeled_paragraph(doc, "软件分类", "应用软件")
    add_labeled_paragraph(doc, "软件说明", "原创")
    add_labeled_paragraph(doc, "开发方式", "单独开发")
    add_labeled_paragraph(doc, "开发完成日期", "2026年3月3日")
    add_labeled_paragraph(doc, "发表状态", "未发表")

    doc.add_heading("二、系统运行与开发环境", level=1)
    add_labeled_paragraph(
        doc,
        "开发的硬件环境",
        "64位台式机或笔记本电脑，四核及以上处理器，16GB内存，512GB SSD。",
    )
    add_labeled_paragraph(
        doc,
        "运行的硬件环境",
        "64位台式机或笔记本电脑，双核及以上处理器，8GB内存，200MB以上可用磁盘空间。",
    )
    add_labeled_paragraph(
        doc,
        "开发该软件的操作系统",
        "Microsoft Windows 11 家庭中文版 64位（10.0.26200）。",
    )
    add_labeled_paragraph(
        doc,
        "软件开发环境/开发工具",
        "Python 3.13.5、Visual Studio Code、Streamlit、GeoPandas、Pandas、NumPy、SciPy、scikit-image、Statsmodels、ReportLab。",
    )
    add_labeled_paragraph(
        doc,
        "软件运行平台/操作系统",
        "Windows 10 / Windows 11 64位。",
    )
    add_labeled_paragraph(
        doc,
        "软件运行支撑环境/支持软件",
        "Python 3.11及以上，Streamlit、GeoPandas、Pandas、Shapely、SciPy、scikit-image、Statsmodels、Folium、Jinja2、ReportLab。",
    )
    add_labeled_paragraph(doc, "编程语言", "Python")
    add_labeled_paragraph(doc, "源程序量", f"约 {sum(count_lines(path) for path in spec.code_files)} 行核心代码。")

    doc.add_heading("三、开发目的与软件定位", level=1)
    add_body_paragraph(
        doc,
        spec.purpose,
    )
    add_body_paragraph(
        doc,
        f"本软件在当前项目中承担的定位是：{spec.positioning}。在申报口径上，该软件作为独立的功能单元进行描述，强调自身在业务边界、核心能力和成果输出上的独立性，以便与同源项目中的其他子模块形成清晰区分。",
    )

    doc.add_heading("四、面向领域与用户对象", level=1)
    add_body_paragraph(doc, spec.audience)

    doc.add_heading("五、软件的主要功能", level=1)
    for paragraph in spec.main_functions:
        add_body_paragraph(doc, paragraph)

    doc.add_heading("六、软件的技术特点", level=1)
    for paragraph in spec.technical_features:
        add_body_paragraph(doc, paragraph)

    path = OUTPUT_DIR / f"{spec.key}_{spec.full_name}{spec.version}_申请表.docx"
    doc.save(path)
    return path


def build_operation_manual(spec: SuiteSpec) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc, f"{spec.full_name}{spec.version} 操作手册", f"简称：{spec.short_name}")

    doc.add_heading("一、软件概述", level=1)
    add_body_paragraph(
        doc,
        f"{spec.full_name}是依托“松辽流域河道变迁与植被响应原型系统”拆分形成的独立申报软件。该软件重点围绕{spec.positioning}展开，采用本地部署方式运行，能够在普通Windows环境下完成任务配置、分析执行、结果展示与成果输出等操作，适用于科研训练、项目演示、课程设计和材料归档等场景。",
    )
    add_body_paragraph(
        doc,
        f"从功能定位上看，{spec.full_name}并非对原始项目的简单更名，而是对原项目中相关模块进行边界重组后形成的独立软件表达。该软件将与自身定位最相关的算法、界面和输出能力整合为统一流程，以形成可单独申报、可独立运行、可清晰说明的软件成果。",
    )

    doc.add_heading("二、建设目标与适用对象", level=1)
    add_body_paragraph(doc, spec.purpose)
    add_body_paragraph(doc, spec.audience)

    doc.add_heading("三、软件组成与业务流程", level=1)
    add_body_paragraph(
        doc,
        "软件整体采用模块化结构组织。用户首先配置研究区域、分析年份、年份步长及报告选项；系统随后调用内置演示数据或预留的数据接口，执行分析逻辑；在处理完成后，系统将结果分发至图表展示、空间数据导出和文档报告模块，最终形成可用于查看、下载和归档的完整成果集合。",
    )
    add_body_paragraph(
        doc,
        "这种流程设计使用户可以在同一套操作路径中完成“参数输入—自动分析—结果查看—成果导出”的闭环。相比传统的分散式处理方式，软件将配置、计算、展示和文档化工作集中在统一界面内完成，减少了重复整理和跨工具切换造成的时间消耗。",
    )

    doc.add_heading("四、系统运行环境", level=1)
    add_body_paragraph(
        doc,
        "推荐在Windows 10或Windows 11 64位系统中使用本软件。运行前需安装Python 3.11及以上版本，并根据项目依赖清单安装Streamlit、GeoPandas、Pandas、NumPy、SciPy、scikit-image、Statsmodels、Folium、Jinja2与ReportLab等支持组件。安装完成后，用户可通过命令“streamlit run streamlit_app.py”启动本地服务，并在浏览器中访问系统界面。",
    )
    add_body_paragraph(
        doc,
        "在硬件方面，建议使用8GB及以上内存的个人电脑进行运行，以确保图表绘制、空间数据转换和文档导出过程具有较好的响应速度。在教学、答辩或演示场景中，普通实验室电脑即可满足软件的基础运行需求。",
    )

    doc.add_heading("五、主要功能介绍与操作说明", level=1)
    for title, paragraph in spec.operation_sections:
        doc.add_heading(title, level=2)
        add_body_paragraph(doc, paragraph)
        add_body_paragraph(
            doc,
            "在实际操作时，用户仅需按界面提示逐步完成参数设置与查看流程。若输入参数不合法，系统会在前端即时给出提示；若数据源不可用或样本不足，系统会以可读方式反馈问题，避免用户在无效配置下继续执行分析任务。",
        )

    doc.add_heading("六、成果输出与归档方式", level=1)
    add_body_paragraph(
        doc,
        "软件每次运行都会自动创建独立的时间戳输出目录，并生成结构化成果文件。输出内容通常包括年度指标CSV文件、缓冲带统计CSV文件、河道中心线与节点GeoJSON文件、关键图表PNG文件，以及根据设置生成的PDF分析报告。该机制使用户能够清晰区分不同运行批次的结果，便于后续整理、对比和归档。",
    )
    add_body_paragraph(
        doc,
        "对于软件著作权申报而言，这种自动化输出也有助于证明软件具备稳定的结果组织能力和完整的业务闭环。用户无需依赖额外排版工具，即可得到可直接用于展示或提交的标准化成果文件。",
    )

    doc.add_heading("七、版本说明", level=1)
    add_body_paragraph(
        doc,
        f"当前文档对应的软件版本为{spec.version}。本版本主要用于形成本项目拆分后的独立软件申报材料，已经具备明确的软件名称、独立的功能说明、完整的操作流程和相应的代码实现基础。后续如需继续提升申报区分度，可在保持现有结构的基础上进一步扩展功能模块或优化界面呈现。",
    )

    path = OUTPUT_DIR / f"{spec.key}_{spec.full_name}{spec.version}_操作手册.docx"
    doc.save(path)
    return path


def build_code_doc(spec: SuiteSpec) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title(doc, f"{spec.full_name}{spec.version} 代码文档", "用于软件著作权申报的核心代码整理稿")

    total_lines = sum(count_lines(path) for path in spec.code_files)

    doc.add_heading("一、代码文档说明", level=1)
    add_body_paragraph(
        doc,
        f"本代码文档根据当前项目的真实实现内容整理而成，围绕“{spec.full_name}”对应的软件边界，选取与该软件定位最相关的核心源代码进行汇编。当前纳入文档的核心代码文件共{len(spec.code_files)}个，合计约{total_lines}行，覆盖入口程序、配置定义、核心业务逻辑与主要功能模块。",
    )
    add_body_paragraph(
        doc,
        "在软著申报口径中，代码文档的作用是体现软件已经形成较为稳定和可识别的程序结构。因此，本稿保留了真实代码中的模块划分、函数命名与主要逻辑，以便从技术层面直接对应软件说明书中的功能描述。若后续需要按照机构要求调整页数，可在现有基础上继续增加更多代码页或仅保留起止页区间。",
    )

    doc.add_heading("二、纳入文档的代码清单", level=1)
    for path in spec.code_files:
        add_body_paragraph(
            doc,
            f"{path.as_posix()}（约{count_lines(path)}行）",
        )

    doc.add_heading("三、核心代码正文", level=1)
    for path in spec.code_files:
        doc.add_heading(path.as_posix(), level=2)
        add_body_paragraph(
            doc,
            describe_code_file(path),
        )
        append_code(doc, path)

    path = OUTPUT_DIR / f"{spec.key}_{spec.full_name}{spec.version}_代码文档.docx"
    doc.save(path)
    return path


def append_code(doc: Document, path: Path, chunk_size: int = 80) -> None:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    for start in range(0, len(lines), chunk_size):
        chunk = "\n".join(lines[start : start + chunk_size])
        paragraph = doc.add_paragraph(style="CodeBlock")
        run = paragraph.add_run(chunk)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        run.font.size = Pt(8)


def count_lines(path: Path) -> int:
    return len(path.read_text(encoding="utf-8").splitlines())


def describe_code_file(path: Path) -> str:
    name = path.name
    mapping = {
        "streamlit_app.py": "该文件是软件的统一启动入口，负责配置源码路径并调用界面渲染逻辑。",
        "config.py": "该文件定义软件运行涉及的全局配置，包括研究区域、默认年份、空间分辨率、输出目录及阈值参数等基础常量。",
        "models.py": "该文件定义核心数据结构，用于描述分析请求、分析结果与栅格观测对象，是业务数据在模块之间传递的基础。",
        "base.py": "该文件定义遥感数据提供者协议，约束数据源模块需要实现的标准方法，确保后续扩展不同数据提供方式时接口一致。",
        "demo_provider.py": "该文件实现离线演示数据提供逻辑，能够在无外部数据依赖的情况下生成确定性的栅格观测结果和边界数据。",
        "gee_provider.py": "该文件预留真实GEE数据接入接口，并在未配置环境时给出清晰错误提示，用于后续扩展。",
        "preprocess.py": "该文件负责多时相观测的解析、插值、归一化和时间窗口校验，是分析任务开始前的数据预处理模块。",
        "indices.py": "该文件实现NDWI、NDVI和FVC等核心遥感指标的计算逻辑，是水体识别与植被估算的基础。",
        "water.py": "该文件负责水体提取，包括Otsu阈值分割、形态学处理、连通域过滤和水体面积计算。",
        "centerline.py": "该文件负责中心线骨架提取、节点识别、河宽估算、迁移量计算以及坡度均值计算等空间分析功能。",
        "vegetation.py": "该文件负责河岸缓冲区植被统计和FVC变化量计算，用于描述植被对河道变化的响应。",
        "drivers.py": "该文件负责构造回归样本并执行OLS回归分析，用于评估多因子对植被变化的影响。",
        "analysis_service.py": "该文件是业务编排核心，负责串联数据加载、指标计算、空间分析、图表生成、导出和报告创建等流程。",
        "export_service.py": "该文件负责CSV、GeoJSON、文本、JSON和图像结果的落盘，是成果归档模块的基础。",
        "report_service.py": "该文件负责将运行结果整理成PDF文档，形成可直接分发的结构化分析报告。",
        "pages.py": "该文件负责构建Streamlit界面，组织参数输入、结果展示、下载按钮和交互地图等前端逻辑。",
    }
    return mapping.get(name, "该文件为当前软件的组成模块之一，承担对应的程序逻辑实现。")


def build_index_doc(paths: Iterable[Path]) -> None:
    doc = Document()
    configure_doc(doc)
    add_title(doc, "拆分版软著材料目录", "基于当前项目生成的两套软件著作权申报文档")
    add_body_paragraph(
        doc,
        "本目录文件用于说明当前工作区已自动生成的软著材料范围。全部文档均依据现有项目代码与功能边界整理形成，已按“分析引擎”和“监测报告平台”两套口径分别输出申请表、操作手册与代码文档，可直接在Microsoft Word中打开并继续修改。",
    )
    for path in paths:
        add_body_paragraph(doc, path.name)
    doc.save(OUTPUT_DIR / "材料目录_请先看此文件.docx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    suite_a = SuiteSpec(
        key="方案一",
        full_name="松辽流域河道植被遥感分析软件",
        short_name="河道植被分析软件",
        version="V1.0",
        positioning="河道与植被变化分析引擎",
        purpose="本软件的开发目的在于为流域遥感研究提供一套可独立运行的分析引擎，将多时相数据准备、遥感指标计算、河道结构识别、植被覆盖估算和驱动因子评估整合到统一的软件流程中。通过这一软件，使用者可以在不依赖复杂外部系统的前提下完成从数据计算到定量分析的核心工作，从而提升科研建模、教学实验和方案论证的效率。",
        audience="本软件主要面向遥感测绘、水文地理、生态环境监测及高校科研教学等领域，适用于需要开展河道演变分析、植被响应识别与多因子诊断的教师、学生、科研人员和项目成员。",
        main_functions=[
            "软件支持在给定研究区域和时间范围内自动完成多时相观测数据的解析与插值处理，并对输入数据进行统一的归一化和有效性校验，使后续分析建立在稳定一致的数据基础之上。",
            "软件内置NDWI、NDVI和FVC等核心指标计算能力，可对研究区域内的水体与植被状态进行定量刻画，其中水体识别结果可进一步用于河道范围提取，植被指标结果可进一步用于生态响应分析。",
            "软件能够完成基于阈值分割、形态学处理和骨架化算法的河道中心线与节点识别，进一步计算河道面积、中心线长度、平均河宽、迁移量和摆动幅度等关键指标，为河道形态变化提供结构化量化结果。",
            "软件可按河岸缓冲区对植被变化进行分层统计，并构建河道变化与植被变化之间的耦合样本，帮助使用者从空间距离和时间序列两个维度观察河道演变对植被覆盖度的影响。",
            "软件支持构建多因子回归样本，对河宽变化、河道迁移、降水、土地利用强度和坡度等变量进行标准化回归分析，以输出定量化的驱动评估结果。",
        ],
        technical_features=[
            "软件采用模块化Python架构，将数据模型、数据源、预处理、指标计算、河道提取、植被统计、回归分析与结果导出分别封装为独立模块，结构清晰，便于维护与扩展。",
            "在算法实现层面，软件综合使用Otsu阈值分割、形态学开闭运算、骨架化、距离变换、缓冲距离统计和OLS回归等技术路径，使河道提取与植被响应分析能够以统一流程完成。",
            "软件支持离线演示数据快速运行，同时保留GEE数据提供接口，为后续接入真实遥感平台预留扩展能力。这种设计既保证当前版本可演示、可申报，又为后续升级保留技术空间。",
            "软件以结构化结果为导向，可输出指标表、空间数据和图形文件，为后续成果展示、论文整理和平台集成提供标准化中间结果。",
        ],
        operation_sections=[
            ("5.1 参数配置与分析启动", "用户进入软件后，可在首页设置研究区域、开始年份、结束年份、年份步长、数据源类型及是否生成报告。软件会首先校验输入参数是否合法，确认年份范围在允许区间内、年份顺序正确且缓冲区参数符合预设逻辑。完成设置后，用户点击“开始分析”，系统即进入自动计算阶段。"),
            ("5.2 多时相数据准备", "软件会根据用户选择的年份范围，优先加载与目标年份匹配的观测数据；若目标年份位于两个基准年份之间，软件将自动执行线性插值，生成对应年份的近似观测结果。该处理方式保证了分析时间轴的连续性，同时使演示环境下的多时相分析具备可运行性和可解释性。"),
            ("5.3 河道提取与结构识别", "在完成基础指标计算后，软件基于NDWI结果进行水体识别，并通过形态学处理清理噪声区域，再计算河道中心线与关键节点。系统最终将以图表和指标表形式展示每个年份的河道面积、中心线长度、平均河宽及迁移信息，使用户能够从结构层面理解河道演变过程。"),
            ("5.4 植被响应分析", "软件根据河道水体结果计算河岸距离，并按预设缓冲区对NDVI和FVC进行分区统计。用户可在结果界面中查看不同缓冲带的植被覆盖变化，并结合年度对比结果识别近河区与远河区在响应强度上的差异。"),
            ("5.5 驱动因子评估", "软件将河道指标、植被变化与辅助驱动变量组合成回归样本，自动执行标准化回归分析并输出系数、显著性和拟合度指标。该功能可帮助用户从统计角度理解不同因素对植被变化的影响方向和相对强度。"),
            ("5.6 结果导出", "当分析完成后，软件会自动生成CSV、GeoJSON、PNG和PDF等文件。用户可将这些结果用于课程汇报、论文支撑材料、内部讨论或其他后续研究工作。"),
        ],
        code_files=[
            ROOT / "streamlit_app.py",
            ROOT / "src" / "river_insight" / "config.py",
            ROOT / "src" / "river_insight" / "domain" / "models.py",
            ROOT / "src" / "river_insight" / "providers" / "base.py",
            ROOT / "src" / "river_insight" / "providers" / "demo_provider.py",
            ROOT / "src" / "river_insight" / "providers" / "gee_provider.py",
            ROOT / "src" / "river_insight" / "pipeline" / "preprocess.py",
            ROOT / "src" / "river_insight" / "pipeline" / "indices.py",
            ROOT / "src" / "river_insight" / "pipeline" / "water.py",
            ROOT / "src" / "river_insight" / "pipeline" / "centerline.py",
            ROOT / "src" / "river_insight" / "pipeline" / "vegetation.py",
            ROOT / "src" / "river_insight" / "pipeline" / "drivers.py",
            ROOT / "src" / "river_insight" / "services" / "analysis_service.py",
        ],
    )

    suite_b = SuiteSpec(
        key="方案二",
        full_name="流域生态监测成果展示与自动报告平台",
        short_name="生态监测报告平台",
        version="V1.0",
        positioning="成果展示、任务编排与自动报告平台",
        purpose="本软件的开发目的在于将流域生态分析过程中的参数配置、结果展示、成果导出和报告生成能力整合成一套平台化工具，使使用者能够在同一界面内完成任务发起、结果浏览、文件下载与材料归档。该软件突出的是面向交付与呈现的应用价值，而不是单纯的底层分析算法能力。",
        audience="本软件主要面向生态监测展示、科研成果汇报、教学演示和项目交付场景，适用于需要快速查看分析结果、整理归档文件并生成正式汇报材料的教师、学生、研究团队和项目负责人。",
        main_functions=[
            "软件支持以本地Web界面的方式统一管理分析任务参数，用户可在一个入口页面中完成研究区域、年份范围、数据源类型和报告选项的设置，避免使用者在多个脚本或工具之间来回切换。",
            "软件能够将分析结果按河道变化、植被响应、驱动评估和报告导出四个页面分区展示，使结构化指标、图表结果和空间结果能够在同一任务上下文中被集中查看。",
            "软件支持交互式地图展示与静态图表混合呈现方式。用户既可以通过图像快速理解总体趋势，也可以通过地图查看边界、中心线与节点等空间对象的分布情况。",
            "软件可将一次分析任务自动整理为带时间戳的输出目录，并生成CSV、GeoJSON、PNG和PDF等多种文件格式，满足汇报、存档、共享与二次加工等多类需求。",
            "软件内置报告生成能力，能够将参数摘要、年度指标、图表与回归结果组合为统一格式的文档成果，减少人工整理材料的时间成本。",
        ],
        technical_features=[
            "软件采用“界面层—分析服务层—导出服务层—报告服务层”的分层结构，界面负责交互，分析服务负责编排，导出服务负责文件落盘，报告服务负责PDF生成，职责边界清晰。",
            "软件使用Streamlit构建本地Web交互界面，结合Folium展示空间对象、Matplotlib输出静态图像，使软件既具备交互性，也能兼顾传统申报和汇报材料所需要的静态文件输出。",
            "软件将单次任务的所有产物统一收口到独立目录中，便于结果留痕和批次管理。该设计对于教学演示、科研归档和项目过程管理都具有较高的实用性。",
            "软件虽然依托同源项目形成，但在申报口径上突出的是成果呈现和自动文档化能力，因此在功能边界上可与偏分析引擎的软件形成明确区分。",
        ],
        operation_sections=[
            ("5.1 平台首页与任务创建", "用户启动系统后，首先进入平台首页。在该页面中，用户可完成研究区域、分析起止年份、年份步长、数据源和是否生成报告等参数设置。软件会实时校验参数合法性，并在必要时给出提示，确保后续任务能够在有效条件下执行。"),
            ("5.2 河道变化页面", "在“河道变化”页面中，平台会展示年度水体掩膜图、年度指标趋势图、中心线叠加图以及交互式地图结果。该页面的核心作用是帮助用户快速掌握分析对象在不同年份下的河道结构变化情况，并通过图形与表格进行综合对照。"),
            ("5.3 植被响应页面", "在“植被响应”页面中，平台展示缓冲带FVC热力图和河道-植被耦合散点图，同时配套输出缓冲带统计表。用户可以在该页面集中查看河岸距离变化与植被覆盖变化之间的关系，并为后续研究结论提供直观支撑。"),
            ("5.4 驱动评估页面", "在“驱动评估”页面中，平台展示标准化回归系数图、样本统计量和系数明细。该页面的设计目标是将复杂的统计结果以更适合演示和汇报的形式呈现出来，使非技术背景的阅读者也能快速理解关键因素的影响方向。"),
            ("5.5 报告导出页面", "在“报告导出”页面中，平台集中列出本次任务生成的全部文件，用户可以直接下载PDF报告、CSV指标文件和GeoJSON空间数据。该页面承担了成果分发与归档出口的作用，是整个平台闭环中与最终交付最直接相关的部分。"),
            ("5.6 自动报告生成与归档", "当用户勾选生成报告后，平台会在分析结束时自动将任务参数、年度指标、图表结果和回归摘要整理为PDF文档，并将所有产物放入按时间戳命名的目录中。该机制尤其适合需要反复运行、快速产出和批量归档的项目场景。"),
        ],
        code_files=[
            ROOT / "streamlit_app.py",
            ROOT / "src" / "river_insight" / "config.py",
            ROOT / "src" / "river_insight" / "domain" / "models.py",
            ROOT / "src" / "river_insight" / "providers" / "demo_provider.py",
            ROOT / "src" / "river_insight" / "services" / "analysis_service.py",
            ROOT / "src" / "river_insight" / "services" / "export_service.py",
            ROOT / "src" / "river_insight" / "services" / "report_service.py",
            ROOT / "src" / "river_insight" / "ui" / "pages.py",
            ROOT / "src" / "river_insight" / "pipeline" / "indices.py",
            ROOT / "src" / "river_insight" / "pipeline" / "water.py",
            ROOT / "src" / "river_insight" / "pipeline" / "centerline.py",
        ],
    )

    generated: list[Path] = []
    for spec in (suite_a, suite_b):
        generated.append(build_application_doc(spec))
        generated.append(build_operation_manual(spec))
        generated.append(build_code_doc(spec))

    build_index_doc(generated)


if __name__ == "__main__":
    main()
